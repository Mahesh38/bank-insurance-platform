#!/usr/bin/env python3
"""Generate the shareable Word work-order for AU Bank DevOps.

Source of truth for agents remains:
  docs/platform/engineering/GITLAB-BANK-DEVOPS-PROVISIONING.md

This script emits a .docx the application team can email / attach without
asking DevOps to read the repository markdown.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x12, 0x33, 0x5A)
GOLD = RGBColor(0xB5, 0x86, 0x2D)
RED = RGBColor(0x8B, 0x1E, 0x1E)
SLATE = RGBColor(0x33, 0x33, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT = "F4F7FA"
HEADER_BG = "12335A"

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUT = ROOT / "docs/platform/engineering/AU-SFB-NIP-GitLab-DevOps-Work-Order.docx"


def shade(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), hex_color)
    fill.set(qn("w:val"), "clear")
    tcPr.append(fill)


def set_cell_text(cell, text: str, *, bold=False, color=None, size=10, center=False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run.bold = bold
    run.font.color.rgb = color or SLATE


def add_table(doc: Document, headers: list[str], rows: list[list[str]], *, col_widths=None) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, HEADER_BG)
        set_cell_text(cell, h, bold=True, color=WHITE, size=9)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            if r_i % 2 == 1:
                shade(cell, ROW_ALT)
            set_cell_text(cell, val, size=9)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"


def body(doc: Document, text: str, *, bold=False, color=None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = "Calibri"
    run.bold = bold
    run.font.color.rgb = color or SLATE


def bullets(doc: Document, items: list[str], *, numbered=False) -> None:
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(item, style=style)
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            run.font.size = Pt(11)
            run.font.name = "Calibri"
            run.font.color.rgb = SLATE


def callout(doc: Document, title: str, text: str, *, fill="F8EFE6") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(title + "  ")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RED
    r.font.name = "Calibri"
    r2 = p.add_run(text)
    r2.font.size = Pt(11)
    r2.font.name = "Calibri"
    r2.font.color.rgb = SLATE
    doc.add_paragraph()


def work_package(doc: Document, code: str, title: str, depth: str, you_do: list[str], evidence: list[str]) -> None:
    heading(doc, f"{code}  —  {title}", 2)
    body(doc, f"Depth of this package: {depth}", bold=True)
    body(doc, "Bank DevOps does:")
    bullets(doc, you_do, numbered=True)
    body(doc, "Evidence to attach to the handover ticket:")
    bullets(doc, evidence)


def set_run_font(run, *, size=11, bold=False, color=None, name="Calibri") -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    run.font.color.rgb = color or SLATE


def header_footer(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run("AU Small Finance Bank  ·  Insurance Distribution Platform (NIP)")
    set_run_font(r, size=9, bold=True, color=NAVY)
    r2 = hp.add_run("\tINTERNAL  ·  GitLab / CI / CD Work Order  ·  Wave 0")
    set_run_font(r2, size=9, color=GOLD)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fr = fp.add_run("NIP-WO-GL-001  ·  Version 1.0  ·  07 September 2026  ·  Application delivery team  ·  Page ")
    set_run_font(fr, size=8, color=SLATE)
    # PAGE field
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run = fp.add_run()
    run._r.append(fld)
    run2 = fp.add_run()
    run2._r.append(instr)
    run3 = fp.add_run()
    run3._r.append(fld2)
    for r in (run, run2, run3):
        set_run_font(r, size=8, color=SLATE)


def cover(doc: Document) -> None:
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AU SMALL FINANCE BANK")
    set_run_font(r, size=14, bold=True, color=GOLD)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Insurance Distribution Platform")
    set_run_font(r, size=22, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("New Insurance Platform (NIP)")
    set_run_font(r, size=14, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("GitLab, CI and CD Provisioning Work Order")
    set_run_font(r, size=18, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("For Bank DevOps  ·  GitLab Administration  ·  Platform Engineering")
    set_run_font(r, size=12, color=SLATE)

    doc.add_paragraph()
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Document ID", "NIP-WO-GL-001"],
            ["Version", "1.0"],
            ["Date", "07 September 2026"],
            ["Classification", "Internal — bank staff with a named role"],
            ["From", "NIP Application Delivery Team (Engineering, SRE, Security, QA, Architecture)"],
            ["To", "AU Bank DevOps assignee — create GitLab, CI templates, CD loc, AWS/IaC"],
            ["Application team must not", "Create GitLab projects, apply Terragrunt/Terraform, use AWS console, own CD"],
            ["Canonical markdown (repo)", "docs/platform/engineering/GITLAB-BANK-DEVOPS-PROVISIONING.md"],
        ],
    )
    callout(
        doc,
        "How to use this file",
        "This is the work order. Execute the work packages in order. Tick the acceptance "
        "sheet at the end. Do not copy any unofficial repository. Create empty, bank-owned "
        "projects, then grant the application team Developer access on application projects only. "
        "They will commit source after the projects exist.",
    )


def build(path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = SLATE
    header_footer(doc)
    cover(doc)

    heading(doc, "1.  Purpose and the split of work", 1)
    body(
        doc,
        "AU Bank policy for this programme: the application (dev / tech) team codes and owns "
        "CI job definitions via reusable templates. Bank DevOps takes this requirement and "
        "builds the GitLab estate, runners, registries, Terragrunt/Terraform, AWS and the "
        "lines of code for CD. The application team will not be GitLab Owners, will not be "
        "AWS console operators, and will not be given Terraform state access.",
    )
    add_table(
        doc,
        ["Surface", "Bank DevOps", "Application team"],
        [
            ["Create GitLab groups, subgroups, projects", "Does", "Does not"],
            ["Users, SSO mapping, roles, protected branches, runners", "Does", "Does not"],
            ["Terragrunt / Terraform / AWS accounts / networking", "Does", "Raises a requirement only — never applies"],
            ["CD (dev → sit → uat → prod → dr)", "Does", "Consumes promoted artefacts only"],
            ["Application source (Java, Flutter, tests, docs)", "Empty project + permissions", "Commits after the repo exists"],
            ["CI templates (build / test / coverage / SAST hooks)", "Owns the paved-road templates", "Includes templates; may extend allowed build/test jobs only"],
            ["Deploy secrets and cloud credentials", "Owns in the approved secret store", "Never in git"],
            ["Container and Maven registries", "Provisions", "Push via CI job identity, not personal keys"],
        ],
    )
    callout(
        doc,
        "Hard rule — loc of CD",
        "Application repositories must not contain Terraform, Terragrunt, AWS account modules, "
        "or a job that applies infrastructure. Those live only under the infra/ subgroup. "
        "A pipeline that can terraform apply from an application project is a defect.",
    )

    heading(doc, "2.  What you are being asked to stand up (scope of depth)", 1)
    body(
        doc,
        "Do Wave 0 completely. Wave 1 empty shells may be created now if that is cheaper "
        "in bank process; do not treat them as live services until shared libraries publish "
        "to the Maven registry.",
    )
    add_table(
        doc,
        ["Wave", "You create", "Why this depth"],
        [
            [
                "Wave 0 — now",
                "1 top-level group, 11 people-access groups, 8 live projects with CI, "
                "branch protection, runners, Maven + container registries, 4 infra projects (empty for you to fill)",
                "Shared Java libraries compile with the services. Splitting every microservice "
                "before a package registry exists will break the build.",
            ],
            [
                "Wave 1 — after Maven registry",
                "Up to 27 extract targets (6 libs + persistence + IAM + domain + integration + platform). "
                "Optional empty shells now.",
                "Each bounded context becomes its own GitLab project. Services depend on released coordinates.",
            ],
            [
                "Not now",
                "Customer BFF CI/CD, customer Flutter, signed Play/App Store upload, DAST as an MR gate, Argo CD",
                "Customer channel is R1. Store binaries are CD + Security. DAST is nightly after an environment exists.",
            ],
        ],
    )

    heading(doc, "3.  Work packages — execute in this order", 1)
    body(
        doc,
        "Each package is one completion outcome. Do not skip WP-04 (CI templates) or WP-05 "
        "(branch protection). A complete group tree with no blocking pipeline is not accepted.",
    )

    work_package(
        doc,
        "WP-01",
        "Top-level GitLab group and subgroup tree",
        "1 top-level group, Internal visibility, plus the subgroups below. No public projects, snippets or packages.",
        [
            "Create top-level group au-bank-insurance-platform (place under the bank parent group if naming policy requires it; keep relative paths).",
            "Create subgroups: platform, backend, frontend, infra, platform-common, ws2-iam, ws3-domain, ws3-integration, ws3-platform, ws3-edge.",
            "Set visibility Internal (or bank equivalent of ‘named role, never public’).",
            "Disable public forks and public packages at group level.",
        ],
        [
            "Screenshot or export of the group tree.",
            "Confirmation that the group is not public.",
        ],
    )
    body(doc, "Required tree:")
    body(
        doc,
        "au-bank-insurance-platform/\n"
        "  platform/          ci-templates, nip-governance\n"
        "  backend/           nip-backend          (Wave 0 Gradle monorepo)\n"
        "  frontend/          nip-app              (Flutter NIP-APP)\n"
        "  infra/             terraform-live, terraform-modules, gitlab-cd, ansible-ops\n"
        "  platform-common/   6 bank-common-* libs + bank-persistence-service\n"
        "  ws2-iam/           identity-provider-adapter-service, identity-authorization-service, workforce-access-bff\n"
        "  ws3-domain/        customer, lead, consent, suitability, product-catalogue, journey-orchestration,\n"
        "                     quotation, proposal, payment, policy-issuance\n"
        "  ws3-integration/   integration-hub-service, 1sb-integration-service, direct-insurer-adapter-service\n"
        "  ws3-platform/      audit-compliance, notification, reporting-mis, administration-config\n"
        "  ws3-edge/          customer-bff (empty — do not wire CI/CD until R1)",
        bold=False,
    )

    work_package(
        doc,
        "WP-02",
        "People access groups, SSO and roles",
        "11 access groups mapped to bank AD/SSO. Named identities only. No shared admin user. "
        "Application team is never Owner of the top group and never Maintainer of infra/*.",
        [
            "Create the access groups in the matrix below and map SSO/AD identities the application team will name in the ticket.",
            "Grant roles exactly as the matrix. Do not ‘helpfully’ give Developer on infra.",
            "Create and rotate deploy/project access tokens yourself, scoped to registry push from CI job identity.",
            "Protect CD environments: prod and dr = nip-devops-admins only, with change-record approval. Application Developer must not include Deploy to production.",
        ],
        [
            "Access-group membership export.",
            "Proof a named application engineer can clone nip-backend and cannot push to infra/terraform-live.",
        ],
    )
    add_table(
        doc,
        ["Access group", "Who", "App projects", "ci-templates", "infra/*"],
        [
            ["nip-devops-admins", "Bank DevOps / GitLab admins", "Maintainer", "Owner", "Owner"],
            ["nip-sre-platform", "Named application SRE + bank platform", "Developer", "Maintainer (MR)", "Reporter"],
            ["nip-engineering", "Application developers", "Developer", "Developer (MR only)", "Reporter"],
            ["nip-tech-leads", "Named tech leads", "Maintainer if bank allows; else Developer + CODEOWNERS", "Developer", "Reporter"],
            ["nip-architecture", "Architecture", "Developer", "Reporter", "Reporter"],
            ["nip-security", "Security", "Developer (read findings)", "Reporter", "Reporter"],
            ["nip-qa", "QA", "Developer or Reporter + artefacts", "Reporter", "Reporter"],
            ["nip-compliance", "Compliance / Risk", "Reporter", "Reporter", "Reporter"],
            ["nip-product", "Product", "Reporter", "Reporter", "None"],
            ["nip-dba", "Database architect", "Developer on persistence only", "Reporter", "Reporter"],
            ["nip-auditors", "Internal audit / SOC", "Reporter", "Reporter", "Reporter"],
        ],
    )
    add_table(
        doc,
        ["Environment", "Who may deploy", "Approval"],
        [
            ["dev", "nip-devops-admins (+ optional nip-sre-platform)", "None or one DevOps"],
            ["sit", "nip-devops-admins", "One DevOps"],
            ["uat", "nip-devops-admins", "DevOps + named tech lead"],
            ["prod", "nip-devops-admins only", "DevOps + change record; Security may block"],
            ["dr", "nip-devops-admins only", "Same as prod"],
        ],
    )

    work_package(
        doc,
        "WP-03",
        "Wave 0 live projects (create these with a default branch and CI on day one)",
        "8 projects. Four are application-facing (templates, backend, frontend, governance). "
        "Four are infra (you fill). Do not seed application business source.",
        [
            "Create the eight projects with default branch main, visibility Internal.",
            "Enable Merge requests, CI/CD, Container Registry (services), Package Registry (Maven) on backend now.",
            "Disable public packages and any cleanup policy that deletes SHA tags younger than 90 days.",
            "Optional DevOps first commit: a one-file .gitlab-ci.yml that only includes the matching template. No application source.",
        ],
        ["Project settings export for each of the eight projects."],
    )
    add_table(
        doc,
        ["Project", "Purpose", "CI include", "Application team role"],
        [
            ["platform/ci-templates", "Paved-road GitLab CI YAML", "n/a — this is the include source", "Developer (MR only)"],
            ["backend/nip-backend", "Java 21 Gradle monorepo: all libs + all services until Wave 1", "java-monorepo.yml + security-sast.yml", "Developer"],
            ["frontend/nip-app", "One Flutter NIP-APP (web + later store binaries). RM / ISR / admin / ops are roles, not extra apps", "flutter-app.yml + security-sast.yml", "Developer"],
            ["platform/nip-governance", "Programme docs and operating model (not a wiki)", "governance-docs.yml", "Developer"],
            ["infra/terraform-live", "Terragrunt live: dev/, sit/, uat/, prod/, dr/. AWS India only (ap-south-1; DR ap-south-2 when S09 says so)", "iac-scan.yml + CD", "Reporter"],
            ["infra/terraform-modules", "Reusable Terraform modules; no live state", "iac-scan.yml", "Reporter"],
            ["infra/gitlab-cd", "Promote an immutable image digest → environment → smoke → release evidence", "CD loc", "Reporter"],
            ["infra/ansible-ops", "Ansible DR drills and post-deploy sanity (bank standard)", "ops pipelines", "Reporter"],
        ],
    )
    body(
        doc,
        "Backend layout the application team will commit later: libs/bank-common-* "
        "(error, domain, security, audit, observability, secrets), services/ (one Gradle "
        "module per bounded context), templates/microservice-skeleton/, Gradle wrapper. "
        "Stack: Java 21, Spring Boot, Gradle Kotlin DSL. Local ports 8080–8084 implemented, "
        "8090–8105 skeletons — do not collide these in deploy manifests.",
    )
    body(
        doc,
        "Frontend: one client, NIP-APP. It talks only to the token-hiding BFF "
        "(workforce-access-bff). It never calls 1SB, never calls a database, never receives "
        "OAuth access or refresh tokens. Wave 0 CI builds web and unsigned artefacts only. "
        "Signed IPA/APK and store upload are CD + Security, not developer self-serve.",
    )

    work_package(
        doc,
        "WP-04",
        "CI templates (paved road)",
        "8 template files. Gate jobs have no allow_failure. Application templates must not "
        "define terraform, terragrunt, aws, or kubectl apply. MR pipelines interruptible: true. "
        "Secret-scan jobs use full git depth.",
        [
            "Commit the template files listed below into platform/ci-templates (names may follow bank convention; the jobs are mandatory).",
            "Wire include: from each Wave 0 application project to the matching files on ref main.",
            "Fail the template review if an application template grows a deploy/apply job.",
        ],
        ["Repository tree of ci-templates and a green include from a sample project."],
    )
    add_table(
        doc,
        ["Template file", "Used by", "Mandatory jobs / commands"],
        [
            ["/templates/java-monorepo.yml", "nip-backend", "./gradlew --no-daemon build test jacocoTestReport jacocoTestCoverageVerification  + ArchUnit summary + artefact publish.  ./gradlew test alone is NOT the gate."],
            ["/templates/java-library.yml", "Wave 1 bank-common-*", "build, test, coverage, publish Maven package"],
            ["/templates/java-service.yml", "Wave 1 services", "build, test, coverage, container build, image scan, SBOM, push digest"],
            ["/templates/flutter-app.yml", "nip-app", "flutter pub get; flutter analyze (errors fail); flutter test --coverage; flutter build web"],
            ["/templates/governance-docs.yml", "nip-governance", "JDK 21 FreshnessCheck (exit 1 = warn, ≥2 = fail); python3 scripts/governance/ci-checks.py; python3 scripts/context/validate-context.py"],
            ["/templates/security-sast.yml", "all application projects", "secret scan, SAST, SCA, SBOM"],
            ["/templates/container-scan.yml", "any image pipeline", "image scan BEFORE registry publish"],
            ["/templates/iac-scan.yml", "infra/* only", "tfsec / Checkov or bank equivalent; never included from application repos"],
        ],
    )
    body(doc, "Java merge-request pipeline (blocks at each gate):")
    bullets(
        doc,
        [
            "secret scan",
            "SAST — new critical / high fails",
            "SCA — critical / high with a fix available fails",
            "compile Java 21",
            "unit + component tests",
            "JaCoCo verification — below threshold fails",
            "ArchUnit — violation fails",
            "static quality — Spotless/Checkstyle or Sonar quality gate",
            "assemble artefacts",
            "SBOM CycloneDX — must list Maven coordinates (a Dart-only SBOM on a Java repo fails)",
            "container build (services) → image scan (critical fails) → publish immutable SHA/digest",
        ],
        numbered=True,
    )
    body(
        doc,
        "Deploy, Terraform apply, AWS smoke and DAST against a live URL are not this pipeline. "
        "Those are WP-09. CI artefact retention: tests 14 days, security 30 days, SBOM 90 days. "
        "JaCoCo paths: **/build/reports/jacoco/test/html/index.html and jacocoTestReport.xml.",
    )
    add_table(
        doc,
        ["CI performance (you own runner capacity)", "Target"],
        [
            ["p95 merge-request feedback", "Under 10 minutes"],
            ["Flake", "Under 1%"],
            ["If the runner pool cannot meet 10 minutes", "That is a DevOps capacity defect — do not drop gates"],
            ["Caching", "Gradle ~/.gradle/caches, Pub, Maven local; parallel jobs"],
            ["Runner tags", "nip-java, nip-flutter, nip-docs, nip-docker"],
            ["Runner images", "Linux; JDK 21; Flutter stable matching sdk ^3.5.4; Python 3.12; DinD or Kaniko for images"],
        ],
    )

    work_package(
        doc,
        "WP-05",
        "Protect main and pin required job names",
        "Branch protection is what makes CI a gate. A skipped required job must not count as passed.",
        [
            "On main (and on uat/prod branches if you use them): no direct push; merge request required; ≥1 approving review; CODEOWNERS for *; dismiss stale approvals; branches up to date; no force-push; no deleting main.",
            "Turn on Pipelines must succeed / Merge when pipeline succeeds. Maintainers on application projects still cannot skip the pipeline.",
            "Pin the exact required job names below. Do not rename them in a way that leaves protection pointing at a missing job.",
            "Trunk-based: short-lived branches, MR to main. If bank policy mandates a long-lived develop, it has the same required jobs as main.",
        ],
        [
            "Branch-protection export showing required job names.",
            "A demonstration MR that fails a unit test and cannot merge.",
        ],
    )
    add_table(
        doc,
        ["Kind", "Required job names to pin"],
        [
            ["Backend / Java", "java:test-and-coverage · java:archunit · security:secret-scan · security:sast · security:sca · security:sbom"],
            ["Frontend / Flutter", "flutter:analyze · flutter:test · security:secret-scan · security:sast · security:sca"],
            ["Governance", "governance:freshness · governance:ci-checks"],
        ],
    )
    body(doc, "Seed CODEOWNERS (application team will maintain content; you create the file):")
    add_table(
        doc,
        ["Path", "Approvers"],
        [
            ["*", "@au-bank-insurance-platform/engineering-reviewers"],
            ["infra/ or **/*.tf", "@au-bank-insurance-platform/devops-only — these paths must not exist in application repos"],
            ["**/adapter/onesb/**", "engineering + architecture"],
            ["**/src/main/**/persistence/**", "engineering + database"],
            ["docs/governance/**", "architecture + delivery"],
        ],
    )
    body(doc, "Optional first-commit include for a Java service project:")
    body(
        doc,
        "include:\n"
        "  - project: 'au-bank-insurance-platform/platform/ci-templates'\n"
        "    file: '/templates/java-service.yml'\n"
        "    ref: main\n"
        "  - project: 'au-bank-insurance-platform/platform/ci-templates'\n"
        "    file: '/templates/security-sast.yml'\n"
        "    ref: main",
    )

    work_package(
        doc,
        "WP-06",
        "Registries, variables and secrets",
        "Maven registry for shared libs; container registry promoting digests; no secrets in application CI variables.",
        [
            "Enable GitLab Package Registry (Maven) for bank-common-* publish from nip-backend.",
            "Enable Container Registry (GitLab or ECR). Promote by digest. Scan-on-push. Never use latest as the promotion mechanism.",
            "Application CI variables: non-secret only (for example SPRING_PROFILES_ACTIVE for test, Sonar host URL). No AWS keys, no DB passwords, no 1SB credentials.",
            "Runtime secrets in AWS Secrets Manager (or bank equivalent), injected at deploy time by CD. You rotate them with Security.",
        ],
        ["Registry URLs and a sample digest tag; variable inventory with no secret values."],
    )

    work_package(
        doc,
        "WP-07",
        "Security scanning in the pipeline",
        "Secret scan, SAST, SCA, image scan, SBOM on every relevant pipeline. Tools may be GitLab Ultimate, "
        "SonarQube (bank standard), gitleaks, Trivy, or the bank equivalent — the fail-closed behaviour is not optional.",
        [
            "Secret scanning: GitLab Secret Detection and/or pinned gitleaks (or bank equivalent that fails the job). Working tree finding blocks the MR. Weekly full-history scan: a historical finding means rotate the credential, not delete the line.",
            "SAST: bank SonarQube (preferred) or GitLab SAST/Semgrep. Java + Dart/Flutter. New critical or high on the MR fails. Existing findings go to the risk register; they do not get allow_failure.",
            "SCA: GitLab Dependency Scanning, Trivy, or bank-approved SCA. Must resolve transitive Java dependencies (lockfiles / built jars), not only build.gradle.kts text. Flutter scans pubspec.lock. Every MR and weekly. Introducing a known critical CVE must fail.",
            "Container image scan before publish. Critical OS or library CVE blocks publish. Non-root, minimal/trusted base, no privileged mode.",
            "SBOM CycloneDX (or bank format), 90-day retention, copy next to the image if that is the bank pattern. Zero Maven coordinates on a Java pipeline = failed job. Generate SBOM even if Legal has not yet issued a licence-fail list.",
            "DAST: nightly against dev once an environment exists. Not a merge-request gate in Wave 0. Application team does not configure DAST.",
            "IaC scan (tfsec/Checkov or bank equivalent) on every infra MR. Public exposure, unencrypted store, over-broad IAM = fail.",
        ],
        [
            "One pipeline run showing secret, SAST, SCA, SBOM (and image scan if a Dockerfile is present).",
            "Proof that a realistic test credential in an MR cannot merge.",
        ],
    )
    add_table(
        doc,
        ["Severity", "Meaning", "Remediation window"],
        [
            ["S0", "Critical, non-bypassable", "Immediate; blocks release"],
            ["S1", "High", "Before the next release"],
            ["S2", "Medium", "Within two releases"],
            ["S3", "Low / hardening", "Backlog"],
        ],
    )
    add_table(
        doc,
        ["Static analysis (quality, not only CVE)", "Tool", "Fail the build when"],
        [
            ["Architecture boundaries", "ArchUnit inside the Java test job", "Any violation (for example 1SB types outside adapter.onesb.*, persistence in the adapter)"],
            ["Formatting / style", "Spotless or Checkstyle", "New violations"],
            ["Code quality", "SonarQube quality gate (bank standard)", "Gate red: new bugs, new vulns, new-code coverage/duplication"],
            ["Flutter", "flutter analyze / dart analyze", "Errors"],
        ],
    )
    body(
        doc,
        "SonarQube does not replace JaCoCo. Both run. JaCoCo is the enforceable module floor; "
        "Sonar is the new-code quality gate. A test step must scan emitted logs for PAN, Aadhaar, "
        "phone, email and health patterns and fail on a match. Do not log OTPs, tokens, passwords "
        "or raw KYC payloads.",
    )

    work_package(
        doc,
        "WP-08",
        "Coverage gates in CI",
        "Fail the build. Do not open a ticket instead. Do not lower lib gates without Tech Lead + QA Lead co-approval and a tracked debt id with expiry.",
        [
            "Wire jacocoTestCoverageVerification into the Java template (already required by WP-04).",
            "Enforce the floors in the table. Exclusions allowed: *Application, package-info, *Config, *Configuration, *Properties.",
            "Flutter: flutter test --coverage; fail if coverage drops below the baseline QA will set. Do not invent a Flutter percentage in this work order.",
        ],
        ["A CI log showing jacocoTestCoverageVerification executed, not merely jacocoTestReport."],
    )
    add_table(
        doc,
        ["Module group", "Line", "Branch"],
        [
            ["libs/*", "≥ 80%", "≥ 70%"],
            ["services:1sb-integration-service", "≥ 90%", "≥ 70%"],
            ["Other services/*", "≥ 50% line (interim floor)", "Not gated yet"],
            ["Compliance-gate code (control paths C1–C10, when those packages exist)", "100% branch", "No waiver"],
        ],
    )

    work_package(
        doc,
        "WP-09",
        "CD loc, Terragrunt/Terraform and AWS (DevOps only)",
        "This is the bulk of your engineering, not the application team’s. They will never apply IaC. "
        "They hand you an image digest + SBOM + scan report + release notes + environment config keys (never secret values).",
        [
            "Fill infra/terraform-live and infra/terraform-modules. Environments: local is theirs; you own dev, sit, uat, prod, dr.",
            "AWS India regions only. No regulated data, backups, logs or archives outside AWS India. Primary ap-south-1; DR ap-south-2 when the S09 pack says so.",
            "Build gitlab-cd to promote an immutable image digest: dev → sit → uat → prod. Never rebuild the image per environment.",
            "Ansible in ansible-ops for DR drills and post-deploy sanity (bank standard).",
            "CloudTrail (who changed AWS) and CloudWatch (how the service runs) are both mandatory. Neither replaces the other.",
            "IaC scan on every infra MR (WP-07). GitLab CI/CD is the CD control plane — do not introduce Argo CD.",
            "Do not connect Render.com or any public PaaS to bank or production-like data.",
        ],
        [
            "Plan or merge request showing Terragrunt live folders for the environments.",
            "CD pipeline definition that consumes a digest, not a rebuild.",
        ],
    )

    work_package(
        doc,
        "WP-10",
        "Repeatable new-project procedure",
        "When Architecture names a new module, you create the GitLab project. The application team does not open GitLab and click New project.",
        [
            "Place the project in the catalogue GitLab group. Project name = Gradle module name (consent-service, not Consent Service).",
            "Default branch main. Visibility Internal. Enable MR, CI/CD, Container Registry (services), Package Registry (libs), and GitLab Secure scans if licensed.",
            "Assign the shared description pointing at platform/ci-templates. Optional include-only first commit (WP-05).",
            "Apply WP-05 protection. Do not wait for the application team to paste business code — they open the first MR with source and tests.",
        ],
        ["Written runbook (this section plus a bank ticket template) owned by DevOps."],
    )

    work_package(
        doc,
        "WP-11",
        "Wave 1 empty shells (optional now, required at extract)",
        "Create empty projects now if bank process is ‘create the shell when architecture names it’. "
        "Otherwise create the groups now (WP-01) and add projects on extract. Do not initialize with sample READMEs that look live.",
        [
            "Optionally create the inventory projects below as empty shells in their groups.",
            "Do not wire DNS, deploy or customer-bff CI.",
        ],
        ["List of created vs deferred Wave 1 projects."],
    )
    add_table(
        doc,
        ["Group", "Project", "Port", "Runtime datastore (not in git)", "Today"],
        [
            ["platform-common", "bank-common-error", "—", "—", "lib, implemented"],
            ["platform-common", "bank-common-domain", "—", "—", "lib, implemented"],
            ["platform-common", "bank-common-security", "—", "—", "lib, implemented"],
            ["platform-common", "bank-common-audit", "—", "—", "lib, implemented"],
            ["platform-common", "bank-common-observability", "—", "—", "lib, implemented"],
            ["platform-common", "bank-common-secrets", "—", "—", "lib, implemented"],
            ["platform-common", "bank-persistence-service", "8081", "Aurora PostgreSQL", "implemented"],
            ["ws2-iam", "identity-provider-adapter-service", "8082", "none", "implemented"],
            ["ws2-iam", "identity-authorization-service", "8083", "Aurora PostgreSQL", "implemented"],
            ["ws2-iam", "workforce-access-bff", "8084", "none (session in cache)", "implemented"],
            ["ws3-domain", "customer-service", "8090", "Aurora PostgreSQL", "skeleton"],
            ["ws3-domain", "lead-service", "8091", "Aurora PostgreSQL", "skeleton"],
            ["ws3-domain", "consent-service", "8092", "Aurora PostgreSQL (append-only)", "skeleton"],
            ["ws3-domain", "suitability-service", "8093", "Aurora PostgreSQL", "skeleton"],
            ["ws3-domain", "product-catalogue-service", "8094", "Aurora + Redis read cache", "skeleton"],
            ["ws3-domain", "journey-orchestration-service", "8095", "DynamoDB state", "skeleton"],
            ["ws3-domain", "quotation-service", "8096", "DynamoDB + Redis", "skeleton"],
            ["ws3-domain", "proposal-service", "8097", "Aurora PostgreSQL", "skeleton"],
            ["ws3-domain", "payment-service", "8098", "Aurora PostgreSQL", "skeleton"],
            ["ws3-domain", "policy-issuance-service", "8099", "Aurora + S3 (PDFs)", "skeleton"],
            ["ws3-integration", "integration-hub-service", "8100", "DynamoDB routing", "skeleton"],
            ["ws3-integration", "1sb-integration-service", "8080", "job store via persistence HTTP", "implemented"],
            ["ws3-integration", "direct-insurer-adapter-service", "8105", "none", "skeleton"],
            ["ws3-platform", "audit-compliance-service", "8101", "DynamoDB + S3 archive", "skeleton"],
            ["ws3-platform", "notification-service", "8102", "DynamoDB delivery log", "skeleton"],
            ["ws3-platform", "reporting-mis-service", "8103", "warehouse over lake", "skeleton"],
            ["ws3-platform", "administration-config-service", "8104", "Aurora PostgreSQL", "skeleton"],
            ["ws3-edge", "customer-bff", "—", "—", "R1 — do not wire"],
            ["frontend", "nip-app", "—", "—", "Wave 0 live"],
        ],
    )
    body(
        doc,
        "Implemented vs skeleton is for your sizing only. Skeleton modules must still compile, "
        "test and pass ArchUnit inside the Wave 0 monorepo. Do not skip CI on them.",
    )

    heading(doc, "4.  What you must not create", 1)
    add_table(
        doc,
        ["Do not", "Why"],
        [
            ["Public GitLab group or public package", "Regulated financial application"],
            ["Application-team AWS console users", "Bank policy — they raise requirements only"],
            ["Terraform inside nip-backend or nip-app", "CD loc belongs in infra/"],
            ["Argo CD", "Bank standard is GitLab CI/CD"],
            ["A second admin/ops Flutter app", "One NIP-APP; roles inside it"],
            ["Customer BFF CI/CD or customer Flutter", "R1"],
            ["Kafka, a second audit database, or Flyway inside 1sb-integration-service", "Standing architecture constraints"],
            ["Environments that store PII outside AWS India", "Residency"],
            ["Render.com or similar attached to bank data", "Never a PII path"],
            ["Shared admin/root GitLab user for the application team", "Named SSO identities"],
        ],
    )

    heading(doc, "5.  Invariants CI must keep enforceable (you do not implement the code)", 1)
    bullets(
        doc,
        [
            "Bank apps never call 1SB or a database directly.",
            "1SB types live only in adapter.onesb.* (ArchUnit).",
            "1sb-integration-service has no Flyway and no JPA. Persistence is platform-common, reached over HTTP.",
            "No PII in logs.",
            "Flutter never receives OAuth tokens; the BFF holds them.",
            "Keycloak is not the source of truth for business authorization.",
        ],
    )

    heading(doc, "6.  Acceptance — Wave 0 is done only when all rows are evidenced", 1)
    body(
        doc,
        "Attach screenshots or exported GitLab settings to the handover ticket. "
        "Until row 4 is demonstrated, engineering foundation remains open regardless of how complete the group tree looks.",
    )
    add_table(
        doc,
        ["#", "Criterion", "Evidence", "Done"],
        [
            ["1", "Top-level group au-bank-insurance-platform exists, Internal, with the WP-01 tree", "Group tree export", "☐"],
            ["2", "Eight Wave 0 projects exist with main protected as WP-05", "Project + protection export", "☐"],
            ["3", "Access groups exist; a named application engineer clones nip-backend and cannot push to infra/terraform-live", "Clone + denied push", "☐"],
            ["4", "ci-templates files exist; a sample MR on nip-backend that fails a unit test cannot merge", "Failed MR screenshot", "☐"],
            ["5", "Required pipeline names are pinned on main; a skipped job is not a green merge", "Protection showing job names", "☐"],
            ["6", "Secret scan, SAST, SCA, SBOM run on Java and Flutter; image scan for Dockerfiles; IaC scan on infra", "Pipeline run", "☐"],
            ["7", "JaCoCo verification is in the Java template with WP-08 floors", "CI log of jacocoTestCoverageVerification", "☐"],
            ["8", "Package + container registries exist; CD promotes digests", "Registry + CD definition", "☐"],
            ["9", "Application projects contain no Terraform and no AWS credentials", "Repo search / template review", "☐"],
            ["10", "Application team has Developer on app projects and a written ‘how to open the first MR’ note (clone URL, runner tags, include path). They still do not have Terraform or AWS.", "Handover note", "☐"],
        ],
    )

    heading(doc, "7.  After Wave 0 (do not do this on day one)", 1)
    bullets(
        doc,
        [
            "Publish bank-common-* to the Maven registry from nip-backend.",
            "For each service, create/activate the Wave 1 project in its group.",
            "Switch that service’s pipeline to java-service.yml.",
            "Leave nip-backend as the umbrella until the last extract.",
            "Split deploy pipelines by group once platform foundation (S09) is green.",
        ],
        numbered=True,
    )

    heading(doc, "8.  Sign-off", 1)
    add_table(
        doc,
        ["Role", "Name", "Date", "Signature"],
        [
            ["Bank DevOps (executor)", "", "", ""],
            ["Application SRE / platform", "", "", ""],
            ["Application Engineering lead", "", "", ""],
            ["Application Security (scanning gates)", "", "", ""],
            ["Application QA (coverage floors)", "", "", ""],
        ],
    )
    body(
        doc,
        "This work order does not authorise production go-live, Terraform apply of the full "
        "S09 landing zone without the architecture pack, or any exception to secret/SAST/SCA gates.",
    )

    heading(doc, "9.  Document control", 1)
    body(
        doc,
        "Canonical machine-readable copy for the application repository: "
        "docs/platform/engineering/GITLAB-BANK-DEVOPS-PROVISIONING.md. If this Word file and "
        "that markdown disagree, update both in the same change. Regenerator: "
        "python3 scripts/platform/generate-gitlab-devops-work-order.py",
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    print(f"wrote {path} ({path.stat().st_size} bytes)")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("-o", "--output", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    build(args.output.resolve())
    return 0


if __name__ == "__main__":
    sys.exit(main())
