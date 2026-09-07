#!/usr/bin/env python3
"""Generate the Phase 1 (max 4 pages) GitLab collaboration brief for bank DevOps.

Full later-phase work-order remains:
  docs/platform/engineering/AU-SFB-NIP-GitLab-DevOps-Work-Order.docx
Markdown source of truth:
  docs/platform/engineering/GITLAB-BANK-DEVOPS-PROVISIONING.md
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x12, 0x33, 0x5A)
GOLD = RGBColor(0xB5, 0x86, 0x2D)
RED = RGBColor(0x8B, 0x1E, 0x1E)
SLATE = RGBColor(0x33, 0x33, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = "12335A"
ROW_ALT = "F4F7FA"
AMBER = "F8EFE6"

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUT = ROOT / "docs/platform/engineering/AU-SFB-NIP-GitLab-Phase-1-Collaboration-Brief.docx"


def shade(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), hex_color)
    fill.set(qn("w:val"), "clear")
    tcPr.append(fill)


def set_run(run, *, size=10, bold=False, color=None) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Calibri"
    run.font.color.rgb = color or SLATE


def cell_text(cell, text: str, *, bold=False, color=None, size=9, center=False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        shade(c, HEADER_BG)
        cell_text(c, h, bold=True, color=WHITE, size=8)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            c = table.rows[r_i + 1].cells[c_i]
            if r_i % 2:
                shade(c, ROW_ALT)
            cell_text(c, val, size=8)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(0)


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(8 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"
        run.font.size = Pt(13 if level == 1 else 11)


def body(doc: Document, text: str, *, bold=False, size=10) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_run(run, size=size, bold=bold)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(0)
        for run in p.runs:
            set_run(run, size=10)


def callout(doc: Document, title: str, text: str, fill=AMBER) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title + "  ")
    set_run(r, size=10, bold=True, color=RED)
    r2 = p.add_run(text)
    set_run(r2, size=10)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def page_field(paragraph) -> None:
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld_begin)
    run2 = paragraph.add_run()
    run2._r.append(instr)
    run3 = paragraph.add_run()
    run3._r.append(fld_end)
    for r in (run, run2, run3):
        set_run(r, size=8, color=SLATE)


def setup(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    r = hp.add_run("AU Small Finance Bank  ·  NIP  ·  GitLab Phase 1 brief")
    set_run(r, size=8, bold=True, color=NAVY)
    r2 = hp.add_run("    INTERNAL  ·  Collaboration start only  ·  4 pages max")
    set_run(r2, size=8, color=GOLD)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fr = fp.add_run("NIP-WO-GL-001-P1  ·  07 Sep 2026  ·  Page ")
    set_run(fr, size=8)
    page_field(fp)
    fr2 = fp.add_run(" of 4  ·  Full later-phase spec is NIP-WO-GL-001 (do not start that yet)")
    set_run(fr2, size=8)


def build(path: Path) -> None:
    doc = Document()
    setup(doc)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(2)
    r = t.add_run("Phase 1 — Open GitLab so the application team can start work")
    set_run(r, size=14, bold=True, color=NAVY)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_after = Pt(6)
    r = s.add_run("Insurance Distribution Platform (NIP)  ·  request to Bank DevOps")
    set_run(r, size=10, color=SLATE)

    add_table(
        doc,
        ["", ""],
        [
            ["Document", "NIP-WO-GL-001-P1  ·  Phase 1 of 3  ·  version 1.0  ·  07 September 2026"],
            ["From / To", "NIP application delivery team  →  AU Bank DevOps assignee"],
            ["Ask", "Create three empty GitLab projects, grant access, protect main, attach a runner. We will commit source via merge requests."],
            ["Do not do in Phase 1", "Terragrunt, Terraform, AWS, CD, extra empty microservice shells, store signing, DAST"],
        ],
    )

    callout(
        doc,
        "Please work in phases.",
        "This four-page brief is the only Phase 1 ask. The long work-order (NIP-WO-GL-001) is the "
        "Phase 2 / Phase 3 specification — keep it for later. Phase 1 is successful when our "
        "developers can clone, branch, open a merge request and see a pipeline run.",
    )

    heading(doc, "1. Three phases — only Phase 1 is requested now", 1)
    add_table(
        doc,
        ["Phase", "Goal", "Bank DevOps does", "Application team does"],
        [
            [
                "1  Now",
                "Start collaboration on source",
                "Group + 3 projects + access + protect main + Java 21 and Flutter runners + skeleton CI",
                "First merge requests with code and tests",
            ],
            [
                "2  Next",
                "Engineering foundation gates",
                "CI templates, secret/SAST/SCA/SBOM, coverage as merge blockers, SonarQube",
                "Keep pipelines green; no Terraform",
            ],
            [
                "3  Later",
                "CD and cloud",
                "Terragrunt/Terraform, AWS India, promote image digests, environments",
                "Hand over digests and config keys only",
            ],
        ],
    )

    heading(doc, "2. Who does what in Phase 1", 1)
    add_table(
        doc,
        ["Action", "Bank DevOps", "Application / dev team"],
        [
            ["Create GitLab group and the three projects", "Yes", "No"],
            ["SSO users, Developer role on the three projects", "Yes", "Send the named list"],
            ["Terragrunt / Terraform / AWS console / CD", "Not in Phase 1", "Will not ask"],
            ["Commit Java, Flutter, docs, tests", "Empty repo only", "Yes — via merge request"],
            ["Skeleton .gitlab-ci.yml (build/test)", "Seed one include file if useful", "We will extend jobs after access exists"],
        ],
    )

    heading(doc, "3. Create exactly this (nothing else)", 1)
    body(
        doc,
        "Top-level group: au-bank-insurance-platform  (Internal — not public). "
        "If bank naming needs a parent, put this group under it and keep the paths below.",
    )
    add_table(
        doc,
        ["Project path", "What it is", "Default branch", "Enable"],
        [
            [
                "backend/nip-backend",
                "Java 21 Gradle monorepo — all shared libs and services in one repo for now",
                "main",
                "MR, CI, Maven package registry, Container registry",
            ],
            [
                "frontend/nip-app",
                "One Flutter app (NIP-APP: web now; store binaries later). RM/admin/ops are roles, not extra apps",
                "main",
                "MR, CI",
            ],
            [
                "platform/nip-governance",
                "Programme docs (process, architecture, SSOT). Not a wiki.",
                "main",
                "MR, CI",
            ],
        ],
    )
    body(
        doc,
        "Do not create the 20+ empty microservice projects, infra/Terraform projects, or customer-BFF in Phase 1. "
        "Those are Phase 2 / 3. Do not copy any unofficial repository — empty projects, we push the first source.",
        bold=True,
    )

    heading(doc, "4. Access so we can collaborate", 1)
    body(doc, "Named SSO identities only. No shared admin user. Send us the clone URLs when done.")
    add_table(
        doc,
        ["GitLab access group", "Members (you map SSO)", "On the 3 projects", "Must not have"],
        [
            ["nip-devops-admins", "Bank DevOps", "Maintainer", "—"],
            ["nip-engineering", "Application developers (list we will attach)", "Developer", "Owner of the top group"],
            ["nip-tech-leads", "Named tech leads", "Developer (or Maintainer if bank policy allows)", "AWS / Terraform rights"],
            ["nip-sre-platform", "Named application SRE", "Developer", "infra Maintainer (infra does not exist yet)"],
        ],
    )

    heading(doc, "5. Protect main — this is how we work", 1)
    bullets(
        doc,
        [
            "No direct push to main. Merge request required. At least one approval.",
            "Pipeline must succeed. A skipped required job must not count as passed.",
            "Dismiss stale approvals on new commits. No force-push, no deleting main.",
            "Trunk-based: short-lived branches, MR into main.",
        ],
    )
    body(doc, "Phase 1 required jobs (pin these names; we will add security jobs in Phase 2):")
    add_table(
        doc,
        ["Project", "Phase 1 required job", "What it must run"],
        [
            ["nip-backend", "java:test-and-coverage", "./gradlew --no-daemon test jacocoTestCoverageVerification   (Java 21)"],
            ["nip-app", "flutter:test", "flutter pub get && flutter analyze && flutter test"],
            ["nip-governance", "governance:ci-checks", "JDK 21 + python3 scripts/governance/ci-checks.py  (we will add the files)"],
        ],
    )
    body(
        doc,
        "Runners: Linux, tag nip-java (Temurin 21 + Gradle cache) and nip-flutter (Flutter SDK ^3.5.4). "
        "If a runner is not ready on day one, still create the projects and access — we can push MRs; "
        "turn the required-job pin on as soon as the runner exists so we do not merge untested code.",
    )

    heading(doc, "6. Optional seed commit (you may do this; we can also do it as MR #1)", 1)
    body(
        doc,
        "A one-file .gitlab-ci.yml per project is enough. No application source. Example for backend:",
    )
    body(
        doc,
        "stages: [test]\n"
        "java:test-and-coverage:\n"
        "  stage: test\n"
        "  tags: [nip-java]\n"
        "  image: eclipse-temurin:21-jdk\n"
        "  script:\n"
        "    - chmod +x gradlew && ./gradlew --no-daemon test jacocoTestCoverageVerification\n"
        "  rules:\n"
        "    - if: $CI_PIPELINE_SOURCE == \"merge_request_event\"\n"
        "    - if: $CI_COMMIT_BRANCH",
        size=8,
    )

    heading(doc, "7. What we will do the day access exists", 1)
    bullets(
        doc,
        [
            "Open MR #1 on nip-backend with the Gradle tree, tests and wrapper.",
            "Open MR #1 on nip-app with the Flutter app.",
            "Open MR #1 on nip-governance with programme docs.",
            "We will not ask you for AWS, Terraform or production deploy in this phase.",
        ],
    )

    heading(doc, "8. Phase 1 acceptance — six ticks, then we start", 1)
    add_table(
        doc,
        ["#", "Done when", "Evidence", "☐"],
        [
            ["1", "Group au-bank-insurance-platform exists, Internal, not public", "Group URL", "☐"],
            ["2", "Three projects exist, empty (or seed CI file only), default branch main", "Three project URLs", "☐"],
            ["3", "Named developers have Developer; they cannot create extra groups", "Access screenshot", "☐"],
            ["4", "main: MR required, ≥1 approval, no direct push", "Protection export", "☐"],
            ["5", "Clone works for one named developer (git clone + create a branch)", "We confirm in the ticket", "☐"],
            ["6", "Runner attached, or a dated plan for nip-java / nip-flutter tags", "Runner list or date", "☐"],
        ],
    )

    heading(doc, "9. Explicitly out of Phase 1", 1)
    body(
        doc,
        "Do not start these until we send the Phase 2 request: CI template project; secret/SAST/SCA/image/SBOM "
        "as merge blockers; SonarQube quality gate; JaCoCo floors as policy (80/70 libs, 90/70 1SB service); "
        "infra/ terraform-live, terraform-modules, gitlab-cd, ansible-ops; AWS accounts; empty shells for "
        "every microservice; customer BFF; Play/App Store signing; Argo CD; public packages.",
    )

    heading(doc, "10. Handover back to us", 1)
    add_table(
        doc,
        ["Please return", "Value"],
        [
            ["Clone URL — nip-backend", ""],
            ["Clone URL — nip-app", ""],
            ["Clone URL — nip-governance", ""],
            ["Who has Developer (names)", ""],
            ["Runner tags live? (yes / date)", ""],
            ["DevOps contact for Phase 2", ""],
        ],
    )
    add_table(
        doc,
        ["Role", "Name", "Date", "Sign"],
        [
            ["Bank DevOps (Phase 1 done)", "", "", ""],
            ["Application tech lead (we can clone)", "", "", ""],
        ],
    )
    body(
        doc,
        "Phase 2 / 3 detail (do not execute now): NIP-WO-GL-001 — "
        "AU-SFB-NIP-GitLab-DevOps-Work-Order.docx. Repo markdown: "
        "docs/platform/engineering/GITLAB-BANK-DEVOPS-PROVISIONING.md.",
        size=9,
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    print(f"wrote {path} ({path.stat().st_size} bytes)")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("-o", "--output", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    build(args.output.resolve())
    return 0


if __name__ == "__main__":
    sys.exit(main())
